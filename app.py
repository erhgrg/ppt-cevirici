import os
from flask import Flask, render_template_string, request, send_file
from pptx import Presentation
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

app = Flask(__name__)
UPLOAD_FOLDER = '/tmp' if os.name != 'nt' else os.path.join(os.path.expanduser("~"), "Desktop", "gecici_donusumler")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="tr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Online PPTX -> PDF Dönüştürücü</title>
    <style>
        body { font-family: 'Arial', sans-serif; background: #f4f7f6; display: flex; justify-content: center; align-items: center; height: 100vh; margin: 0; }
        .container { background: white; padding: 30px; border-radius: 12px; box-shadow: 0 4px 15px rgba(0,0,0,0.1); text-align: center; max-width: 400px; width: 90%; }
        h2 { color: #333; margin-bottom: 20px; }
        input[type="file"] { display: none; }
        .file-label { display: block; background: #0078D4; color: white; padding: 12px; border-radius: 6px; cursor: pointer; font-weight: bold; margin-bottom: 15px; transition: 0.3s; }
        .file-label:hover { background: #005a9e; }
        .btn-submit { background: #107C41; color: white; border: none; padding: 12px 25px; border-radius: 6px; font-weight: bold; cursor: pointer; font-size: 16px; width: 100%; display: none; }
        .btn-submit:hover { background: #0b592e; }
        .status { margin-top: 15px; color: #666; font-size: 14px; }
    </style>
</head>
<body>
    <div class="container">
        <h2>PPTX to PDF Converter</h2>
        <form action="/convert" method="post" enctype="multipart/form-data" onsubmit="showLoading()">
            <label for="file-upload" class="file-label" id="label-text">📂 PowerPoint Dosyası Seç</label>
            <input id="file-upload" type="file" name="file" accept=".pptx, .ppt" onchange="fileSelected()">
            <button type="submit" id="submit-btn" class="btn-submit">🔒 PDF'e Dönüştür ve İndir</button>
        </form>
        <div id="status" class="status"></div>
    </div>
    <script>
        function fileSelected() {
            const input = document.getElementById('file-upload');
            const label = document.getElementById('label-text');
            const btn = document.getElementById('submit-btn');
            if(input.files.length > 0) {
                label.innerText = "✓ Dosya Seçildi";
                label.style.background = "#2b88d8";
                btn.style.display = "block";
            }
        }
        function showLoading() {
            document.getElementById('status').innerText = "Dönüştürülüyor... Lütfen bekleyin...";
            document.getElementById('submit-btn').style.display = "none";
        }
    </script>
</body>
</html>
"""

def linux_pptx_to_pdf(input_path, output_path):
    prs = Presentation(input_path)
    c = canvas.Canvas(output_path, pagesize=letter)
    for slide in prs.slides:
        text_content = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_content.append(shape.text.strip())
        
        y_position = 750
        c.setFont("Helvetica", 12)
        c.drawString(50, 770, "--- Slayt Sayfasi ---")
        for text in text_content:
            c.drawString(50, y_position, text[:80])
            y_position -= 20
            if y_position < 50:
                c.showPage()
                y_position = 750
        c.showPage()
    c.save()

@app.route('/')
def home():
    return render_template_string(HTML_TEMPLATE)

@app.route('/convert', methods=['POST'])
def convert():
    if 'file' not in request.files: return "Dosya yok", 400
    file = request.files['file']
    if file.filename == '': return "Dosya secilmedi", 400
    
    if file:
        input_path = os.path.join(UPLOAD_FOLDER, file.filename)
        file.save(input_path)
        output_filename = os.path.splitext(file.filename)[0] + ".pdf"
        output_path = os.path.join(UPLOAD_FOLDER, output_filename)
        
        try:
            linux_pptx_to_pdf(input_path, output_path)
            return send_file(output_path, as_attachment=True)
        except Exception as e:
            return f"Hata: {str(e)}", 500
        finally:
            if os.path.exists(input_path): os.remove(input_path)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000)
